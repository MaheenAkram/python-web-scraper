"""Create Excel/CSV files only when the GUI requests an export."""

import csv
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLUMNS = [
    ("text", "Extracted Text"),
    ("value", "Price / Value"),
    ("source_url", "Source URL"),
    ("scraped_at", "Scraped At (UTC)"),
]


def validate_export(records, filename, extension):
    if not records:
        raise ValueError("There are no results to export. Scrape a page first.")
    path = Path(filename)
    if path.suffix.lower() != extension:
        raise ValueError(f"Choose a filename ending with {extension}.")
    if not path.parent.is_dir():
        raise ValueError("The selected destination folder does not exist.")
    return path


def cell_text(value):
    text = str(value if value is not None else "")
    return ILLEGAL_CHARACTERS_RE.sub("", text)


def is_web_url(value):
    try:
        parts = urlsplit(value)
        return (
            parts.scheme.lower() in ("http", "https")
            and bool(parts.hostname)
            and parts.username is None
            and parts.password is None
            and not any(character.isspace() for character in value)
        )
    except ValueError:
        return False


def export_excel(records, filename):
    path = validate_export(records, filename, ".xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scraped Data"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="166534")
    alternate_fill = PatternFill("solid", fgColor="F0FDF4")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    thin_side = Side(style="thin", color="DCE9DF")
    borders = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )

    sheet.append([label for key, label in COLUMNS])
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = borders
    sheet.row_dimensions[1].height = 31

    for row_index, record in enumerate(records, start=2):
        for column_index, (key, label) in enumerate(COLUMNS, start=1):
            text = cell_text(record.get(key, ""))
            cell = sheet.cell(row=row_index, column=column_index)
            # Force scraped content to text, preventing Excel formula execution.
            cell.value = text[:32767]
            cell.data_type = "s"
            cell.fill = alternate_fill if row_index % 2 == 0 else white_fill
            cell.font = Font(name="Calibri", size=11, color="1F2937")
            cell.border = borders
            cell.alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )

            if key == "source_url" and is_web_url(text):
                cell.hyperlink = text
                cell.font = Font(
                    name="Calibri", size=11, color="166534", underline="single"
                )

        sheet.row_dimensions[row_index].height = 33

    # Estimate readable widths with a cap for long text and URLs.
    caps = [70, 28, 60, 29]
    for column_index, cap in enumerate(caps, start=1):
        longest = max(
            len(line)
            for row in sheet.iter_rows(
                min_col=column_index, max_col=column_index
            )
            for line in (str(row[0].value or "").splitlines() or [""])
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            cap, max(16, longest + 3)
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    sheet.print_title_rows = "1:1"
    workbook.properties.title = "Web Scraper Results"
    workbook.properties.creator = "Web Scraper"
    try:
        workbook.save(path)
    finally:
        workbook.close()


def safe_csv_value(value):
    text = str(value if value is not None else "")
    # CSV readers may interpret these prefixes as spreadsheet formulas.
    # A leading apostrophe preserves the content as literal text in Excel.
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(
        ("\t", "\r", "\n")
    ):
        return "'" + text
    return text


def export_csv(records, filename):
    path = validate_export(records, filename, ".csv")
    # UTF-8 with BOM makes Unicode easier to open correctly in Excel.
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for key, label in COLUMNS])
        for record in records:
            writer.writerow([
                safe_csv_value(record.get(key, "")) for key, label in COLUMNS
            ])
